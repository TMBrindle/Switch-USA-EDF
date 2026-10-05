"""Pin the EIA tables behind the S0 fuel-price path (s0_production.fuel_prices; CHANGES §67).

Downloads (to --cache, default s0_workflow/data/raw/fuel/, gitignored) EIA's published files. api.eia.gov is not
needed (it is blocked from some environments); these are the same numbers:
  STEO, monthly workbook (the latest edition at download)    https://www.eia.gov/outlooks/steo/xls/STEO_m.xlsx
    Table 2 / 6 / 7a: CLEUDUS, NGEUDUS (cost of coal / natural gas to the electric power sector, nominal $/MMBtu);
    Table 5a NGEPCON (natural gas consumed by the electric power sector, bcf/d); Table 6 CLEPCON_TON (coal,
    million short tons); Table 9a CICPIUS (CPI-U, 1982-84 = 1.00)
  AEO2026 Table 3, Energy Prices by Sector and Source (Electric Power: natural gas, steam coal; 2025 $/MMBtu)
    reference case                    https://www.eia.gov/outlooks/aeo/excel/aeotab3.xlsx
    low / high oil and gas supply     https://www.eia.gov/outlooks/aeo/excel/sidecases/{lowogs,highogs}/aeotab3.xlsx
  AEO2026 supplemental Table 54.1-54.25, Electric Power Projections by EMM Region (reference case): fuel prices to
    the electric power sector (coal, natural gas; 2025 $/MMBtu) and fuel consumption (quadrillion Btu)
                                      https://www.eia.gov/outlooks/aeo/supplement/excel/sup_elec.xlsx

Writes (s0_workflow/data/fuel/):
  steo_power_fuel_monthly.csv      month, cleudus, ngeudus, ngepcon_bcfd, clepcon_mst, cicpius
  aeo2026_power_fuel_national.csv  case, year, fuel, price_2025usd_per_mmbtu
  aeo2026_power_fuel_emm.csv       emm, emm_name, year, fuel, price_2025usd_per_mmbtu, consumption_quads
  SOURCES.yml                      editions, release dates, URLs, retrieval date, sha256 of each download
--check: download again (or use --cache) and compare with the pinned tables; no writes. Exit 1 on a difference.

usage (repo root): python s0_workflow/scripts/fetch_fuel_prices_eia.py [--check] [--offline]
"""
import argparse
import hashlib
import re
import sys
import urllib.request
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
import yaml

REPO = Path(__file__).resolve().parents[2]
OUT = REPO / "s0_workflow/data/fuel"
CACHE = REPO / "s0_workflow/data/raw/fuel"
AEO = "https://www.eia.gov/outlooks/aeo/"
FILES = {
    "steo_m": "https://www.eia.gov/outlooks/steo/xls/STEO_m.xlsx",
    "aeo_ref_tab3": AEO + "excel/aeotab3.xlsx",
    "aeo_lowogs_tab3": AEO + "excel/sidecases/lowogs/aeotab3.xlsx",
    "aeo_highogs_tab3": AEO + "excel/sidecases/highogs/aeotab3.xlsx",
    "aeo_sup_elec": AEO + "supplement/excel/sup_elec.xlsx",
}
CASES = {"aeo_ref_tab3": "reference", "aeo_lowogs_tab3": "low_ogs", "aeo_highogs_tab3": "high_ogs"}
STEO_SERIES = {"CLEUDUS": "cleudus", "NGEUDUS": "ngeudus", "NGEPCON": "ngepcon_bcfd", "CLEPCON_TON": "clepcon_mst",
               "CICPIUS": "cicpius"}


def download(key: str, cache: Path, offline: bool) -> Path:
    cache.mkdir(parents=True, exist_ok=True)
    f = cache / Path(FILES[key]).name.replace(".xlsx", f".{key}.xlsx")
    if not offline:
        req = urllib.request.Request(FILES[key], headers={"User-Agent": "Mozilla/5.0 (Switch-USA-EDF research)"})
        with urllib.request.urlopen(req, timeout=180) as r:
            data = r.read()
        if data[:2] != b"PK":
            raise RuntimeError(f"{FILES[key]}: not an xlsx file (EIA returned a page, e.g. 404)")
        f.write_bytes(data)
    if not f.exists():
        raise FileNotFoundError(f"{f} (run without --offline)")
    return f


def sha256(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


# ---------------------------------------------------------------------------------------------- STEO
def steo(path: Path) -> tuple[pd.DataFrame, dict]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    dates = [r for r in wb["Dates"].iter_rows(max_row=3, values_only=True)]
    edition = str(dates[0][3])                                   # e.g. "September 2026"
    completed = dates[1][3]
    first_year = None
    series = {}
    for sn in wb.sheetnames:
        if not sn.endswith("tab"):
            continue
        for r in wb[sn].iter_rows(values_only=True):
            if r and r[0] == "Forecast date:" and first_year is None:
                first_year = int(r[2])
            if r and r[0] in STEO_SERIES and STEO_SERIES[r[0]] not in series:
                series[STEO_SERIES[r[0]]] = [float(v) if isinstance(v, (int, float)) else np.nan for v in r[2:]]
    missing = set(STEO_SERIES.values()) - set(series)
    if missing:
        raise ValueError(f"STEO {edition}: series not found: {sorted(missing)}")
    n = min(len(v) for v in series.values())
    n = n - n % 12
    months = pd.period_range(f"{first_year}-01", periods=n, freq="M")
    df = pd.DataFrame({"month": months.astype(str), **{k: v[:n] for k, v in series.items()}})
    meta = {"edition": edition, "modeling_completed": str(getattr(completed, "date", lambda: completed)()),
            "months": f"{months[0]}..{months[-1]}"}
    return df, meta


# ---------------------------------------------------------------------------------------------- AEO
def _aeo_meta(d: pd.DataFrame) -> dict:
    meta = {}
    for i in range(10):
        row = [x for x in d.iloc[i].tolist() if isinstance(x, (str, int, float)) and str(x) != "nan"]
        if "Scenario" in row:
            j = row.index("Scenario")
            meta["scenario"], meta["scenario_label"] = str(row[j + 1]), str(row[j + 2]) if len(row) > j + 2 else ""
        if "Release Date" in row:
            meta["release"] = str(row[row.index("Release Date") + 1])
        if "Report" in row:
            meta["report"] = str(row[row.index("Report") + 1])
    return meta


def _year_header(d: pd.DataFrame, start=0) -> tuple[int, dict]:
    for i in range(start, len(d)):
        v = d.iloc[i, 2]
        if isinstance(v, (int, float)) and not pd.isna(v) and 2000 <= v <= 2100:
            return i, {c: int(d.iloc[i, c]) for c in range(2, d.shape[1])
                       if isinstance(d.iloc[i, c], (int, float)) and not pd.isna(d.iloc[i, c])}
    raise ValueError("no year header")


def aeo_national(path: Path, case: str) -> tuple[pd.DataFrame, dict]:
    d = pd.read_excel(path, header=None)
    meta = _aeo_meta(d)
    units = " ".join(str(x) for x in d.iloc[:12, 1] if isinstance(x, str))
    if "2025 dollars" not in units:
        raise ValueError(f"{path.name}: expected 2025 dollars, got {units!r}")
    _, years = _year_header(d)
    ids = d[0].astype(str)
    rows = []
    for key, fuel in (("PRC000:ga_NaturalGas", "naturalgas"), ("PRC000:ga_SteamCoal", "coal")):
        m = ids == key                                           # ga_ = Electric Power, real 2025 dollars
        if m.sum() != 1:
            raise ValueError(f"{path.name}: row {key} not found once")
        r = d[m].iloc[0]
        rows += [{"case": case, "year": y, "fuel": fuel, "price_2025usd_per_mmbtu": float(r[c])} for c, y in years.items()]
    return pd.DataFrame(rows), meta


def aeo_emm(path: Path) -> tuple[pd.DataFrame, dict]:
    d = pd.read_excel(path, header=None)
    meta = _aeo_meta(d)
    col1 = d[1].astype(str)
    titles = d.index[col1.str.match(r"\s*54\.\d+\. Electric Power Projections by Electricity Market Module Region")]
    rows = []
    for t in titles:
        name = str(d.iloc[t + 2, 1]).strip()
        m = re.match(r"(\d+) - (.+)", name)
        code, region = m.group(1), m.group(2)
        _, years = _year_header(d, t)
        end = titles[titles > t].min() if (titles > t).any() else len(d)
        blk = d.iloc[t:end]
        ids = blk[0].astype(str)
        for fuel, price_key, cons_key in (("naturalgas", "sa_NaturalGas", "ra_NaturalGas"), ("coal", "sa_Coal", "ra_Coal")):
            p = blk[ids.str.endswith(price_key)]
            q = blk[ids.str.endswith(cons_key)]
            if len(p) != 1 or len(q) != 1:
                raise ValueError(f"sup_elec: region {code}: {price_key} / {cons_key} not found once")
            unit = blk.loc[: p.index[0], 1].astype(str)
            if "2025 dollars per MMBtu" not in unit.iloc[-3:].str.cat(sep=" ") and "2025 dollars" not in \
                    unit.iloc[-6:].str.cat(sep=" "):
                raise ValueError(f"sup_elec: region {code}: price units not 2025 dollars")
            for c, y in years.items():
                rows.append({"emm": code, "emm_name": region, "year": y, "fuel": fuel,
                             "price_2025usd_per_mmbtu": float(p.iloc[0, c]), "consumption_quads": float(q.iloc[0, c])})
    out = pd.DataFrame(rows)
    if out.emm.nunique() != 25:
        raise ValueError(f"sup_elec: {out.emm.nunique()} EMM regions, expected 25")
    return out, meta


# ---------------------------------------------------------------------------------------------- main
def build(cache: Path, offline: bool):
    paths = {k: download(k, cache, offline) for k in FILES}
    s, smeta = steo(paths["steo_m"])
    nat, metas = [], {}
    for k, case in CASES.items():
        df, metas[case] = aeo_national(paths[k], case)
        nat.append(df)
    emm, emeta = aeo_emm(paths["aeo_sup_elec"])
    sources = {
        "retrieved": str(date.today()),
        "steo": {"edition": smeta["edition"], "modeling_completed": smeta["modeling_completed"],
                 "months": smeta["months"], "url": FILES["steo_m"], "sha256": sha256(paths["steo_m"]),
                 "series": STEO_SERIES, "units": "prices nominal $/MMBtu; NGEPCON bcf/d; CLEPCON_TON million short "
                                                  "tons; CICPIUS index 1982-84 = 1.00"},
        "aeo2026": {k: {**metas[CASES[k]], "url": FILES[k], "sha256": sha256(paths[k]),
                        "rows": "Electric Power: Natural Gas (PRC000:ga_NaturalGas), Steam Coal (PRC000:ga_SteamCoal); "
                                "2025 $/MMBtu"} for k in CASES},
        "aeo2026_emm": {**emeta, "url": FILES["aeo_sup_elec"], "sha256": sha256(paths["aeo_sup_elec"]),
                        "tables": "54.1-54.25 Electric Power Projections by EMM Region: Fuel Prices to the Electric "
                                  "Power Sector (2025 $/MMBtu), Fuel Consumption (quadrillion Btu)"},
    }
    return s, pd.concat(nat, ignore_index=True), emm, sources


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--check", action="store_true", help="compare with the pinned tables; no writes")
    ap.add_argument("--offline", action="store_true", help="use the files already in --cache")
    ap.add_argument("--cache", type=Path, default=CACHE)
    a = ap.parse_args(argv)
    s, nat, emm, sources = build(a.cache, a.offline)
    tables = {"steo_power_fuel_monthly.csv": s, "aeo2026_power_fuel_national.csv": nat,
              "aeo2026_power_fuel_emm.csv": emm}
    if a.check:
        bad = []
        for name, df in tables.items():
            pinned = pd.read_csv(OUT / name, dtype={"emm": str})
            new = df.reset_index(drop=True)
            same = list(pinned.columns) == list(new.columns) and len(pinned) == len(new)
            if same:
                for c in new.columns:
                    if pd.api.types.is_numeric_dtype(new[c]):
                        same &= bool(np.allclose(pinned[c].astype(float), new[c].astype(float), rtol=1e-9,
                                                 atol=1e-9, equal_nan=True))
                    else:
                        same &= bool((pinned[c].astype(str) == new[c].astype(str)).all())
            if not same:
                bad.append(name)
        pin = yaml.safe_load(open(OUT / "SOURCES.yml"))
        if pin["steo"]["edition"] != sources["steo"]["edition"]:
            bad.append(f"STEO edition {sources['steo']['edition']} (pinned {pin['steo']['edition']})")
        print("CHECK OK" if not bad else "CHECK DIFFERS: " + "; ".join(bad))
        return 1 if bad else 0
    OUT.mkdir(parents=True, exist_ok=True)
    for name, df in tables.items():
        df.to_csv(OUT / name, index=False, float_format="%.10g")
    with open(OUT / "SOURCES.yml", "w") as f:
        yaml.safe_dump(sources, f, sort_keys=False)
    print(f"STEO {sources['steo']['edition']}; AEO2026 {sources['aeo2026']['aeo_ref_tab3'].get('release')}; "
          f"wrote {', '.join(tables)} and SOURCES.yml to {OUT.relative_to(REPO)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
